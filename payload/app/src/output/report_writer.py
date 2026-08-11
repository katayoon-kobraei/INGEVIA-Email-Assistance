from __future__ import annotations

import csv
import os
import re
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

REPORT_LOG_FILENAME = "email_report_log.csv"
REPORT_XLSX_FILENAME = "Informe de Emails.xlsx"
REPORT_FIELDNAMES = [
    "EntryID", "Date", "Time", "Sender Name", "Sender Email",
    "Subject", "Summary", "Path",
]

_WHITESPACE_RE = re.compile(r"\s+")


def cheap_fallback_summary(email, max_chars=140):
    body = email.get("body") or ""
    cleaned = _WHITESPACE_RE.sub(" ", body).strip()
    if len(cleaned) <= max_chars:
        return cleaned
    return cleaned[:max_chars].rstrip() + "..."


def _report_log_path(output_root):
    return os.path.join(output_root, REPORT_LOG_FILENAME)


def _upgrade_report_schema(path: Path) -> None:
    if not path.exists():
        return
    try:
        with path.open("r", newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle)
            if (reader.fieldnames or []) == REPORT_FIELDNAMES:
                return
            rows = list(reader)
    except (OSError, csv.Error):
        return

    fd, temp_name = tempfile.mkstemp(prefix="email_report_", suffix=".csv", dir=str(path.parent))
    os.close(fd)
    temp_path = Path(temp_name)
    try:
        with temp_path.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(handle, fieldnames=REPORT_FIELDNAMES)
            writer.writeheader()
            for row in rows:
                writer.writerow({name: row.get(name, "") for name in REPORT_FIELDNAMES})
        os.replace(temp_path, path)
    finally:
        try:
            temp_path.unlink(missing_ok=True)
        except OSError:
            pass


def append_to_report_log(email, sender_label, summary, folder_path, output_root):
    path = Path(_report_log_path(output_root))
    path.parent.mkdir(parents=True, exist_ok=True)
    _upgrade_report_schema(path)
    file_exists = path.exists() and path.stat().st_size > 0
    sender_email = email.get("sender_email") or email.get("sender") or email.get("recipient") or ""
    sender_name = email.get("sender_name") or sender_label or ""

    with path.open("a", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=REPORT_FIELDNAMES)
        if not file_exists:
            writer.writeheader()
        writer.writerow({
            "EntryID": email.get("id") or "",
            "Date": email["timestamp"].strftime("%Y-%m-%d"),
            "Time": email["timestamp"].strftime("%H:%M"),
            "Sender Name": sender_name,
            "Sender Email": sender_email,
            "Subject": email["subject"],
            "Summary": summary or "",
            "Path": folder_path,
        })


def _load_priority_maps(output_root):
    priority_path = os.path.join(output_root, "priorities.csv")
    by_id = {}
    by_key = {}
    if not os.path.exists(priority_path):
        return by_id, by_key
    try:
        with open(priority_path, newline="", encoding="utf-8-sig") as f:
            for row in csv.DictReader(f):
                try:
                    value = int(row.get("Priority") or row.get("Prioridad") or 0)
                except (TypeError, ValueError):
                    continue
                if not 1 <= value <= 5:
                    continue
                entry_id = str(row.get("EntryID") or "").strip()
                date = str(row.get("Date") or row.get("Fecha") or "").strip()
                subject = " ".join(str(row.get("Subject") or row.get("Asunto") or "").split()).casefold()
                if entry_id:
                    by_id[entry_id] = value
                by_key[(date, subject)] = value
    except (OSError, csv.Error):
        pass
    return by_id, by_key


def generate_email_report_xlsx(output_root):
    log_path = Path(_report_log_path(output_root))
    if not log_path.exists():
        return
    _upgrade_report_schema(log_path)

    with log_path.open(newline="", encoding="utf-8-sig") as f:
        data_rows = list(csv.DictReader(f))
    if not data_rows:
        return

    by_id, by_key = _load_priority_maps(output_root)
    display_header = [
        "Date", "Time", "Sender Name", "Sender Email", "Subject",
        "Priority", "Summary", "Path",
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = "Emails procesados"
    ws.append(display_header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"

    row_fills = {
        5: PatternFill("solid", fgColor="F8D7DA"),
        4: PatternFill("solid", fgColor="FFE0B2"),
        3: PatternFill("solid", fgColor="FFF3CD"),
    }
    priority_fills = {
        5: PatternFill("solid", fgColor="D32F2F"),
        4: PatternFill("solid", fgColor="F57C00"),
        3: PatternFill("solid", fgColor="FBC02D"),
    }

    for source in data_rows:
        date_text = " ".join(
            part for part in (str(source.get("Date") or ""), str(source.get("Time") or "")) if part
        ).strip()
        subject_key = " ".join(str(source.get("Subject") or "").split()).casefold()
        entry_id = str(source.get("EntryID") or "").strip()
        priority = by_id.get(entry_id) if entry_id else None
        if priority is None:
            priority = by_key.get((date_text, subject_key), "")

        ws.append([
            source.get("Date", ""), source.get("Time", ""),
            source.get("Sender Name", ""), source.get("Sender Email", ""),
            source.get("Subject", ""), priority, source.get("Summary", ""),
            source.get("Path", ""),
        ])
        row_number = ws.max_row
        if priority in row_fills:
            for cell in ws[row_number]:
                cell.fill = row_fills[priority]
            priority_cell = ws.cell(row=row_number, column=6)
            priority_cell.fill = priority_fills[priority]
            priority_cell.font = Font(
                bold=True,
                color="FFFFFF" if priority in (4, 5) else "000000",
            )

    widths = [12, 8, 22, 28, 34, 12, 50, 60]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    summary_col_index = display_header.index("Summary") + 1
    for row in ws.iter_rows(min_row=2, min_col=summary_col_index, max_col=summary_col_index):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.auto_filter.ref = ws.dimensions
    out_path = os.path.join(output_root, REPORT_XLSX_FILENAME)
    try:
        wb.save(out_path)
    except PermissionError:
        print(f"  (Couldn't update {REPORT_XLSX_FILENAME} -- it's probably open in Excel. Will retry next run.)")
