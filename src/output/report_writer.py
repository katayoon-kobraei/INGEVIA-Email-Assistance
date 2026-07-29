import csv
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

REPORT_LOG_FILENAME = "email_report_log.csv"
REPORT_XLSX_FILENAME = "Informe de Emails.xlsx"

_WHITESPACE_RE = re.compile(r"\s+")


def cheap_fallback_summary(email, max_chars=140):
    """Used only when no LLM call happened at all for this email -- the
    Plenergy deterministic US-code fast-path match is the one place in
    the pipeline that files an email without ever calling Gemini, so
    there's no model output to pull a summary from. This takes the
    first stretch of the body instead. Lower quality than the
    LLM-written summaries used everywhere else, but free, and better
    than leaving the report cell blank."""
    body = email.get("body") or ""
    cleaned = _WHITESPACE_RE.sub(" ", body).strip()
    if len(cleaned) <= max_chars:
        return cleaned
    return cleaned[:max_chars].rstrip() + "..."


def _report_log_path(output_root):
    return os.path.join(output_root, REPORT_LOG_FILENAME)


def append_to_report_log(email, sender_label, summary, folder_path, output_root):
    """Durable, append-only source of truth for the tracking report --
    one row per email that actually got saved somewhere (project,
    Plenergy, or billing). Cheap and crash-safe; the .xlsx report is
    just a regenerated view of this file, rebuilt fresh at the end of
    every run -- so a failed .xlsx save never loses data."""
    path = _report_log_path(output_root)
    file_exists = os.path.exists(path)
    sender_email = email.get("sender") or email.get("recipient") or ""

    with open(path, "a", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        if not file_exists:
            writer.writerow(["Date", "Time", "Sender Name", "Sender Email", "Subject", "Summary", "Path"])
        writer.writerow([
            email["timestamp"].strftime("%Y-%m-%d"),
            email["timestamp"].strftime("%H:%M"),
            sender_label or "",
            sender_email,
            email["subject"],
            summary or "",
            folder_path,
        ])


def generate_email_report_xlsx(output_root):
    """Rebuilds the human-facing .xlsx report from the CSV log above.
    Regenerated fully every run rather than appended to directly,
    because openpyxl can't cheaply append one row to an existing
    workbook without reloading and re-saving the whole file anyway --
    and rebuilding from the CSV means a failed save (e.g. someone has
    the report open in Excel) never loses data, it just retries clean
    on the next run."""
    log_path = _report_log_path(output_root)
    if not os.path.exists(log_path):
        return  # nothing filed yet -- nothing to report

    with open(log_path, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))

    if not rows:
        return

    header, data_rows = rows[0], rows[1:]

    wb = Workbook()
    ws = wb.active
    ws.title = "Emails procesados"

    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"

    for row in data_rows:
        ws.append(row)

    # Reasonable column widths + a wrapped Summary column so the sheet
    # is readable without anyone having to manually resize it.
    widths = [12, 8, 22, 28, 34, 50, 60]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    summary_col_index = header.index("Summary") + 1 if "Summary" in header else 6
    for row in ws.iter_rows(min_row=2, min_col=summary_col_index, max_col=summary_col_index):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.auto_filter.ref = ws.dimensions

    out_path = os.path.join(output_root, REPORT_XLSX_FILENAME)
    try:
        wb.save(out_path)
    except PermissionError:
        # Almost always means someone has the report open in Excel --
        # the CSV log above already has everything durably saved, so
        # this just quietly retries clean on the next run.
        print(f"  (Couldn't update {REPORT_XLSX_FILENAME} -- it's probably open in Excel. Will retry next run.)")